"""엔드투엔드 검증 스크립트.

pytest 없이도 그대로 실행할 수 있다::

    python tests/test_end_to_end.py

회귀·분류·EDA 세 경로와 예외 상황(전부 결측, 상수 컬럼, 한글 cp949 인코딩,
Excel 입력)을 모두 통과해야 성공으로 간주한다.
"""

from __future__ import annotations

import sys
import tempfile
import warnings
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from autoeda import AnalysisConfig, load_dataframe, run_analysis, run_from_file  # noqa: E402
from autoeda.profiling import (  # noqa: E402
    CATEGORICAL,
    NUMERIC,
    TEXT,
    infer_column_type,
    profile_dataframe,
    suggest_task_type,
)

SAMPLES = ROOT / "samples"
PASSED: list[str] = []
FAILED: list[str] = []


def check(name: str, condition: bool, detail: str = "") -> None:
    (PASSED if condition else FAILED).append(f"{name}{' — ' + detail if detail else ''}")
    print(f"  {'✅' if condition else '❌'} {name}{' — ' + detail if detail else ''}")


# ---------------------------------------------------------------------------
def test_type_inference() -> None:
    print("\n[1] 컬럼 타입 추론")
    check("수치형 인식", infer_column_type(pd.Series(np.random.randn(200))) == NUMERIC)
    check(
        "저카디널리티 정수 → 범주형",
        infer_column_type(pd.Series([1, 2, 3] * 50)) == CATEGORICAL,
    )
    check(
        "문자 범주형 인식",
        infer_column_type(pd.Series(["서울", "부산", "대구"] * 40)) == CATEGORICAL,
    )
    check(
        "식별자 → 텍스트",
        infer_column_type(pd.Series([f"ID{i:05d}" for i in range(300)])) == TEXT,
    )
    check(
        "천단위 콤마 문자열 → 수치형",
        infer_column_type(pd.Series([f"{i * 1000:,}" for i in range(200)])) == NUMERIC,
    )
    check(
        "날짜 문자열 인식",
        infer_column_type(
            pd.Series(pd.date_range("2024-01-01", periods=200).astype(str))
        ) == "datetime",
    )


def test_task_detection() -> None:
    print("\n[2] 문제 유형 자동 판정")
    df = pd.DataFrame(
        {
            "price": np.random.gamma(2, 1000, 300),
            "grade": np.random.choice(["A", "B", "C"], 300),
            "flag": np.random.choice([0, 1], 300),
        }
    )
    check("연속 수치 → 회귀", suggest_task_type(df, "price") == "regression")
    check("문자 라벨 → 분류", suggest_task_type(df, "grade") == "classification")
    check("이진 정수 → 분류", suggest_task_type(df, "flag") == "classification")
    check("목표 없음 → EDA", suggest_task_type(df, None) == "eda")


def test_regression_pipeline(out_dir: Path) -> None:
    print("\n[3] 회귀 전체 파이프라인")
    path = SAMPLES / "sample_sales_regression.csv"
    if not path.exists():
        check("샘플 파일 존재", False, str(path))
        return

    config = AnalysisConfig(
        target="매출액", analysis_type="auto", cv_folds=3, max_models=5,
        report_title="회귀 검증 보고서",
    )
    result = run_from_file(path, config, output_dir=out_dir)

    check("회귀로 판정", result.task_type == "regression")
    check("모델 생성", result.modeling is not None)
    if result.modeling:
        r2 = result.modeling.holdout_metrics.get("R2", -9)
        check("R² > 0.5", r2 > 0.5, f"R²={r2:.3f}")
        check("리더보드 비어있지 않음", not result.modeling.leaderboard.empty)
        check("변수 중요도 산출", not result.modeling.feature_importance.empty)
    check("차트 생성", len(result.charts) >= 4, f"{len(result.charts)}장")
    check(
        "보고서 파일 생성",
        result.report_path is not None and result.report_path.exists()
        and result.report_path.stat().st_size > 50_000,
    )


def test_classification_pipeline(out_dir: Path) -> None:
    print("\n[4] 분류 전체 파이프라인")
    path = SAMPLES / "sample_quality_classification.csv"
    if not path.exists():
        check("샘플 파일 존재", False, str(path))
        return

    config = AnalysisConfig(
        target="품질판정", analysis_type="auto", cv_folds=3, max_models=5,
        exclude_columns=["로트번호"], report_title="분류 검증 보고서",
    )
    result = run_from_file(path, config, output_dir=out_dir)

    check("분류로 판정", result.task_type == "classification")
    check("모델 생성", result.modeling is not None)
    if result.modeling:
        acc = result.modeling.holdout_metrics.get("정확도", 0)
        check("정확도 > 0.6", acc > 0.6, f"정확도={acc:.3f}")
        check("혼동 행렬 생성", result.modeling.confusion is not None)
        check("ROC-AUC 산출", "ROC-AUC" in result.modeling.holdout_metrics)
        check("클래스 라벨 보존", result.modeling.class_labels == ["불량", "양품"])
    check("보고서 파일 생성", result.report_path is not None and result.report_path.exists())


def test_eda_only(out_dir: Path) -> None:
    print("\n[5] 목표 변수 없는 EDA")
    df = pd.read_csv(SAMPLES / "sample_sales_regression.csv")
    result = run_analysis(
        df, AnalysisConfig(report_title="EDA 검증 보고서"),
        source_meta={"파일명": "sample.csv"}, output_dir=out_dir,
    )
    check("EDA 로 처리", result.task_type == "eda")
    check("모델 없음", result.modeling is None)
    check("보고서 생성", result.report_path is not None and result.report_path.exists())


def test_edge_cases(out_dir: Path) -> None:
    print("\n[6] 예외 상황 처리")
    rng = np.random.default_rng(0)
    n = 220
    df = pd.DataFrame(
        {
            "전부결측": [np.nan] * n,
            "상수컬럼": ["동일값"] * n,
            "식별번호": [f"NO{i:06d}" for i in range(n)],
            "구분": rng.choice(["가", "나", "다"], n),
            "값1": rng.normal(50, 10, n),
            "값2": rng.normal(0, 1, n),
        }
    )
    df["결과"] = df["값1"] * 2 + df["값2"] * 5 + rng.normal(0, 3, n)
    df.loc[rng.choice(n, 20, replace=False), "값1"] = np.nan

    profile = profile_dataframe(df)
    check("전부결측 컬럼 제외 권고", "전부결측" in profile.suggested_exclusions)
    check("상수 컬럼 제외 권고", "상수컬럼" in profile.suggested_exclusions)
    check("식별자 컬럼 제외 권고", "식별번호" in profile.suggested_exclusions)

    result = run_analysis(
        df, AnalysisConfig(target="결과", cv_folds=3, max_models=4),
        source_meta={"파일명": "edge.csv"}, output_dir=out_dir,
    )
    check("결측·상수 포함 데이터 분석 성공", result.modeling is not None)
    if result.modeling:
        check(
            "높은 설명력 확보",
            result.modeling.holdout_metrics.get("R2", 0) > 0.7,
            f"R²={result.modeling.holdout_metrics.get('R2', 0):.3f}",
        )

    # 매우 작은 데이터
    tiny = pd.DataFrame({"x": range(30), "y": [i * 2 + (i % 3) for i in range(30)]})
    small = run_analysis(
        tiny, AnalysisConfig(target="y", cv_folds=2, max_models=3),
        source_meta={"파일명": "tiny.csv"}, output_dir=out_dir, build_docx=False,
    )
    check("소량 데이터도 처리", small.modeling is not None)


def test_io_formats(out_dir: Path) -> None:
    print("\n[7] 입출력 형식")
    df = pd.DataFrame(
        {
            "지역": ["서울", "부산", "대구", "광주"] * 30,
            "금액": np.random.default_rng(1).integers(1000, 9000, 120),
        }
    )

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)

        cp949 = tmp_path / "cp949.csv"
        df.to_csv(cp949, index=False, encoding="cp949")
        loaded = load_dataframe(cp949)
        check("cp949 CSV 읽기", list(loaded.df.columns) == ["지역", "금액"], loaded.encoding or "")

        semicolon = tmp_path / "semi.csv"
        df.to_csv(semicolon, index=False, sep=";", encoding="utf-8-sig")
        loaded = load_dataframe(semicolon)
        check("세미콜론 구분자 인식", loaded.df.shape[1] == 2)

        excel = tmp_path / "data.xlsx"
        df.to_excel(excel, index=False)
        loaded = load_dataframe(excel)
        check("Excel 읽기", loaded.df.shape == df.shape)

        try:
            load_dataframe(tmp_path / "없는파일.csv")
            check("없는 파일 예외 발생", False)
        except Exception:
            check("없는 파일 예외 발생", True)


def test_report_structure(out_dir: Path) -> None:
    print("\n[8] 보고서 구조 검증")
    from docx import Document

    reports = sorted(out_dir.glob("*.docx"), key=lambda p: p.stat().st_mtime)
    if not reports:
        check("보고서 존재", False)
        return

    doc = Document(str(reports[0]))
    headings = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    text = "\n".join(p.text for p in doc.paragraphs)

    check("제목 스타일 문단 존재", len(headings) >= 5, f"{len(headings)}개")
    check("표 삽입됨", len(doc.tables) >= 4, f"{len(doc.tables)}개")
    check("이미지 삽입됨", len(doc.inline_shapes) >= 3, f"{len(doc.inline_shapes)}장")
    check("분석 요약 섹션", "분석 요약" in text)
    check("결론 섹션", "결론" in text)


def main() -> int:
    warnings.filterwarnings("ignore")
    print("=" * 64)
    print("Standalone Auto-EDA & Report Builder — 엔드투엔드 검증")
    print("=" * 64)

    with tempfile.TemporaryDirectory() as tmp:
        out_dir = Path(tmp) / "reports"
        out_dir.mkdir(parents=True, exist_ok=True)

        test_type_inference()
        test_task_detection()
        test_regression_pipeline(out_dir)
        test_classification_pipeline(out_dir)
        test_eda_only(out_dir)
        test_edge_cases(out_dir)
        test_io_formats(out_dir)
        test_report_structure(out_dir)

    print("\n" + "=" * 64)
    print(f"통과 {len(PASSED)}건 / 실패 {len(FAILED)}건")
    if FAILED:
        print("\n실패 항목:")
        for item in FAILED:
            print(f"  - {item}")
    print("=" * 64)
    return 1 if FAILED else 0


if __name__ == "__main__":
    raise SystemExit(main())
