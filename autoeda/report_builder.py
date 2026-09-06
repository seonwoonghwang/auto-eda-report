"""Word(.docx) 보고서 자동 빌드 모듈.

python-docx 로 표지 → 목차 → 본문(통계표·차트) → 부록 순의 실무용 레이아웃을
구성한다. 한글 폰트(맑은 고딕)를 문서 스타일과 동아시아 폰트 슬롯에 함께
지정하여 한글이 기본 폰트로 대체되지 않도록 한다.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .automl import ModelingResult
from .config import AnalysisConfig, system_summary
from .insights import (
    correlation_sentences,
    data_summary_sentences,
    executive_summary,
    importance_sentences,
    modeling_sentences,
    recommendation_sentences,
)
from .profiling import DataProfile, top_correlations
from .visualization import ChartSet, chart_index

# 문서 색상 -----------------------------------------------------------------
ACCENT = RGBColor(0x1C, 0x5C, 0xAB)      # 제목/강조 (파랑 550)
INK = RGBColor(0x0B, 0x0B, 0x0B)         # 본문
SUBTLE = RGBColor(0x52, 0x51, 0x4E)      # 캡션/보조
HEADER_FILL = "1C5CAB"                   # 표 머리행 배경
BAND_FILL = "F2F6FC"                     # 표 줄무늬 배경

BODY_FONT = "맑은 고딕"
LATIN_FONT = "Malgun Gothic"

MAX_TABLE_ROWS = 40
CONTENT_WIDTH_CM = 16.0


# ---------------------------------------------------------------------------
# 저수준 헬퍼
# ---------------------------------------------------------------------------
def _set_run_font(run, size: float = 10.5, bold: bool = False,
                  color: RGBColor = INK, font: str = BODY_FONT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), font)


def _configure_styles(doc: Document) -> None:
    """문서 기본 스타일에 한글 폰트를 적용한다."""
    for style_name, size, bold, color in (
        ("Normal", 10.5, False, INK),
        ("Heading 1", 16, True, ACCENT),
        ("Heading 2", 13, True, ACCENT),
        ("Heading 3", 11.5, True, INK),
        ("Caption", 9, False, SUBTLE),
    ):
        try:
            style = doc.styles[style_name]
        except KeyError:  # pragma: no cover - 템플릿에 따라 없을 수 있음
            continue
        style.font.name = LATIN_FONT
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), LATIN_FONT)
        rfonts.set(qn("w:hAnsi"), LATIN_FONT)
        rfonts.set(qn("w:eastAsia"), BODY_FONT)

    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.35


def _shade(cell, color_hex: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color_hex)
    cell._element.get_or_add_tcPr().append(shd)


def _add_field(paragraph, instruction: str) -> None:
    """PAGE, TOC 등 Word 필드를 삽입한다."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, end):
        run._element.append(element)


def _add_page_number_footer(doc: Document) -> None:
    footer = doc.sections[-1].footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(para, "PAGE")
    for run in para.runs:
        _set_run_font(run, size=9, color=SUBTLE)


# ---------------------------------------------------------------------------
# 블록 단위 빌더
# ---------------------------------------------------------------------------
class ReportBuilder:
    """Word 보고서 생성기."""

    def __init__(self, config: AnalysisConfig, template: Optional[Path] = None):
        self.config = config
        self.doc = Document(str(template)) if template and Path(template).exists() else Document()
        if not template:
            self._setup_page()
        _configure_styles(self.doc)

    # -- 페이지/기본 --------------------------------------------------------
    def _setup_page(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)

    def heading(self, text: str, level: int = 1) -> None:
        para = self.doc.add_heading(text, level=level)
        for run in para.runs:
            _set_run_font(
                run,
                size={1: 16, 2: 13, 3: 11.5}.get(level, 11),
                bold=True,
                color=ACCENT if level <= 2 else INK,
            )
        para.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        para.paragraph_format.space_after = Pt(6)

    def para(self, text: str, *, size: float = 10.5, bold: bool = False,
             color: RGBColor = INK, align=None, bullet: bool = False) -> None:
        """본문 문단. `**강조**` 표기를 굵은 글씨로 렌더링한다."""
        paragraph = self.doc.add_paragraph(style="List Bullet" if bullet else None)
        if align is not None:
            paragraph.alignment = align

        for chunk, is_bold in _split_emphasis(text):
            run = paragraph.add_run(chunk)
            _set_run_font(run, size=size, bold=bold or is_bold, color=color)

    def bullets(self, items: Iterable[str]) -> None:
        for item in items:
            self.para(item, bullet=True)

    def caption(self, text: str) -> None:
        if not text:
            return
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(text)
        _set_run_font(run, size=9, color=SUBTLE)
        paragraph.paragraph_format.space_after = Pt(12)

    def spacer(self, points: int = 6) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(points)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # -- 표 ----------------------------------------------------------------
    def table(
        self,
        df: pd.DataFrame,
        *,
        title: Optional[str] = None,
        max_rows: int = MAX_TABLE_ROWS,
        font_size: float = 9,
        numeric_format: str = "{:,.4g}",
    ) -> None:
        if df is None or df.empty:
            self.para("표시할 데이터가 없습니다.", color=SUBTLE, size=9.5)
            return

        if title:
            self.para(title, bold=True, size=10.5)

        truncated = len(df) > max_rows
        view = df.head(max_rows)

        table = self.doc.add_table(rows=1, cols=len(view.columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True

        # 머리행
        for idx, col in enumerate(view.columns):
            cell = table.rows[0].cells[idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(col))
            _set_run_font(run, size=font_size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            _shade(cell, HEADER_FILL)

        # 본문행
        for r_idx, (_, row) in enumerate(view.iterrows()):
            cells = table.add_row().cells
            for c_idx, value in enumerate(row):
                cell = cells[c_idx]
                cell.text = ""
                paragraph = cell.paragraphs[0]
                text, is_number = _format_cell(value, numeric_format)
                run = paragraph.add_run(text)
                _set_run_font(run, size=font_size)
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.RIGHT if is_number else WD_ALIGN_PARAGRAPH.LEFT
                )
                if r_idx % 2 == 1:
                    _shade(cell, BAND_FILL)

        if truncated:
            self.caption(f"※ 전체 {len(df):,}행 중 상위 {max_rows}행만 표시했습니다.")
        else:
            self.spacer(8)

    # -- 이미지 -------------------------------------------------------------
    def image(self, path: Path, caption: str = "", width_cm: float = CONTENT_WIDTH_CM) -> None:
        if not Path(path).exists():
            return
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
        self.caption(caption)

    # -- 저장 ---------------------------------------------------------------
    def save(self, path: Path) -> Path:
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        _add_page_number_footer(self.doc)
        self.doc.save(str(path))
        return path


def _split_emphasis(text: str) -> List[tuple[str, bool]]:
    """`**굵게**` 마크업을 (텍스트, 굵기) 조각으로 분리."""
    parts: List[tuple[str, bool]] = []
    buffer = text
    while "**" in buffer:
        before, _, rest = buffer.partition("**")
        if before:
            parts.append((before, False))
        emphasized, marker, buffer = rest.partition("**")
        if not marker:  # 짝이 맞지 않으면 원문 유지
            parts.append((emphasized, False))
            return parts
        parts.append((emphasized, True))
    if buffer:
        parts.append((buffer, False))
    return parts or [(text, False)]


def _format_cell(value, numeric_format: str) -> tuple[str, bool]:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return "-", False
    if isinstance(value, (int,)) and not isinstance(value, bool):
        return f"{value:,}", True
    if isinstance(value, float):
        if pd.isna(value):
            return "-", False
        if value == int(value) and abs(value) < 1e15:
            return f"{int(value):,}", True
        return numeric_format.format(value), True
    return str(value), False


# ---------------------------------------------------------------------------
# 보고서 조립
# ---------------------------------------------------------------------------
def build_report(
    *,
    output_path: Path,
    config: AnalysisConfig,
    profile: DataProfile,
    raw_preview: pd.DataFrame,
    source_meta: dict,
    preprocessing_table: pd.DataFrame,
    charts: ChartSet,
    modeling: Optional[ModelingResult] = None,
    load_notes: Optional[Sequence[str]] = None,
    template: Optional[Path] = None,
) -> Path:
    """분석 결과 전체를 하나의 Word 문서로 조립한다."""
    rb = ReportBuilder(config, template=template)
    index = chart_index(charts)
    now = datetime.now()

    # ---------- 표지 ----------
    rb.spacer(90)
    title_p = rb.doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(title_p.add_run(config.report_title), size=26, bold=True, color=ACCENT)

    subtitle = config.report_subtitle or f"대상 데이터: {source_meta.get('파일명', '-')}"
    sub_p = rb.doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_run_font(sub_p.add_run(subtitle), size=13, color=SUBTLE)

    rb.spacer(40)
    cover_rows = [
        {"항목": "작성일", "내용": now.strftime("%Y년 %m월 %d일 %H:%M")},
        {"항목": "분석 유형", "내용": _task_label(modeling, config)},
        *[{"항목": k, "내용": str(v)} for k, v in source_meta.items()],
    ]
    if config.organization:
        cover_rows.append({"항목": "소속", "내용": config.organization})
    if config.author:
        cover_rows.append({"항목": "작성자", "내용": config.author})
    rb.table(pd.DataFrame(cover_rows), font_size=10)

    rb.page_break()

    # ---------- 목차 ----------
    rb.heading("목차", level=1)
    toc_p = rb.doc.add_paragraph()
    _add_field(toc_p, r'TOC \o "1-2" \h \z \u')
    rb.caption("※ Word에서 목차 위를 클릭한 뒤 F9 키를 누르면 페이지 번호가 갱신됩니다.")
    rb.page_break()

    # ---------- 1. 분석 요약 ----------
    rb.heading("1. 분석 요약", level=1)
    rb.bullets(executive_summary(profile, source_meta.get("파일명", "데이터"), modeling))

    if charts.font_warning:
        rb.para(f"※ {charts.font_warning}", color=SUBTLE, size=9)

    # ---------- 2. 데이터 개요 ----------
    rb.heading("2. 데이터 개요", level=1)
    for sentence in data_summary_sentences(profile, source_meta.get("파일명", "데이터")):
        rb.para(sentence)
    for note in load_notes or []:
        rb.para(f"※ {note}", color=SUBTLE, size=9.5)

    rb.heading("2.1 데이터 기본 정보", level=2)
    basic = pd.DataFrame(
        [
            {"항목": "행 수", "값": f"{profile.n_rows:,}"},
            {"항목": "열 수", "값": f"{profile.n_cols:,}"},
            {"항목": "메모리 사용량", "값": f"{profile.memory_mb:.2f} MB"},
            {"항목": "중복 행", "값": f"{profile.n_duplicated_rows:,}"},
            {"항목": "수치형 컬럼", "값": f"{len(profile.numeric_columns)}"},
            {"항목": "범주형 컬럼", "값": f"{len(profile.categorical_columns)}"},
            {"항목": "날짜형 컬럼", "값": f"{len(profile.datetime_columns)}"},
        ]
    )
    rb.table(basic, font_size=9.5)

    rb.heading("2.2 컬럼별 요약", level=2)
    rb.table(profile.overview_table(), max_rows=60)

    if config.include_raw_preview and raw_preview is not None and not raw_preview.empty:
        rb.heading("2.3 원본 데이터 미리보기", level=2)
        rb.table(raw_preview.head(10), max_rows=10, font_size=8)

    if "missing" in index:
        rb.heading("2.4 결측값 현황", level=2)
        rb.image(index["missing"].path, index["missing"].caption)

    # ---------- 3. 전처리 내역 ----------
    rb.page_break()
    rb.heading("3. 자동 전처리 내역", level=1)
    rb.para(
        "아래 처리는 사전에 정의된 규칙에 따라 자동 수행되었으며, 모든 단계는 학습·예측 시 "
        "동일하게 재현되도록 하나의 파이프라인으로 구성되어 있습니다."
    )
    rb.table(preprocessing_table, max_rows=60)

    # ---------- 4. 탐색적 데이터 분석 ----------
    rb.page_break()
    rb.heading("4. 탐색적 데이터 분석", level=1)

    describe = profile.numeric_describe_table()
    if not describe.empty:
        rb.heading("4.1 수치형 변수 기술통계", level=2)
        rb.table(describe, max_rows=40)

    categorical = profile.categorical_table()
    if not categorical.empty:
        rb.heading("4.2 범주형 변수 요약", level=2)
        rb.table(categorical, max_rows=30)

    if "distribution" in index:
        rb.heading("4.3 변수 분포", level=2)
        rb.image(index["distribution"].path, index["distribution"].caption)
    if "boxplot" in index:
        rb.image(index["boxplot"].path, index["boxplot"].caption)

    if "correlation" in index or profile.correlation is not None:
        rb.heading("4.4 상관관계 분석", level=2)
        for sentence in correlation_sentences(profile):
            rb.para(sentence)
        if "correlation" in index:
            rb.image(index["correlation"].path, index["correlation"].caption, width_cm=14)
        pairs = top_correlations(profile.correlation, top_n=12)
        if not pairs.empty:
            rb.table(pairs, title="상관계수 상위 변수쌍", max_rows=12)

    if "target" in index:
        rb.heading("4.5 목표 변수 분포", level=2)
        rb.image(index["target"].path, index["target"].caption)

    group_charts = charts.by_key("group_mean")
    if group_charts:
        rb.heading("4.6 주요 구분별 비교", level=2)
        for chart in group_charts:
            rb.image(chart.path, chart.caption)

    # ---------- 5. 예측 모델링 ----------
    if modeling is not None:
        rb.page_break()
        rb.heading("5. 예측 모델링", level=1)
        for sentence in modeling_sentences(modeling):
            rb.para(sentence)

        rb.heading("5.1 후보 모델 비교", level=2)
        rb.table(modeling.leaderboard_display(), max_rows=20)
        if "model_comparison" in index:
            rb.image(index["model_comparison"].path, index["model_comparison"].caption)

        rb.heading("5.2 최적 모델 검증 성능", level=2)
        rb.para(f"선정 모델: **{modeling.best_model_name}**")
        rb.table(modeling.holdout_table(), font_size=9.5)

        if modeling.task_type == "classification":
            if "confusion" in index:
                rb.image(index["confusion"].path, index["confusion"].caption, width_cm=12)
            if "roc" in index:
                rb.image(index["roc"].path, index["roc"].caption, width_cm=11)
        else:
            if "pred_actual" in index:
                rb.image(index["pred_actual"].path, index["pred_actual"].caption, width_cm=12)
            if "residual" in index:
                rb.image(index["residual"].path, index["residual"].caption, width_cm=12)

        # ---------- 6. 변수 중요도 ----------
        rb.page_break()
        rb.heading("6. 주요 영향 변수 (Feature Importance)", level=1)
        for sentence in importance_sentences(modeling):
            rb.para(sentence)
        if "importance" in index:
            rb.image(index["importance"].path, index["importance"].caption)
        if modeling.feature_importance is not None and not modeling.feature_importance.empty:
            rb.table(
                modeling.feature_importance.head(config.top_features),
                title="변수 중요도 상세",
                max_rows=config.top_features,
            )

        if modeling.warnings_:
            rb.heading("6.1 분석 시 유의사항", level=2)
            rb.bullets(modeling.warnings_)

    # ---------- 7. 결론 및 권고 ----------
    rb.page_break()
    rb.heading("7. 결론 및 실무 권고사항", level=1)
    rb.bullets(recommendation_sentences(profile, modeling))

    # ---------- 부록 ----------
    rb.heading("부록. 분석 환경 및 설정", level=1)
    env = system_summary()
    settings = {
        "결측 대체(수치형)": config.numeric_impute,
        "결측 대체(범주형)": config.categorical_impute,
        "이상치 처리": "IQR 윈저라이징" if config.handle_outliers else "미적용",
        "스케일링": config.scaling,
        "검증 데이터 비율": f"{config.test_size:.0%}",
        "교차검증 폴드": str(config.cv_folds),
        "난수 시드": str(config.random_state),
        "차트 폰트": charts.font_name,
    }
    appendix = pd.DataFrame(
        [{"항목": k, "값": v} for k, v in {**env, **settings}.items()]
    )
    rb.table(appendix, font_size=9.5)
    rb.para(
        "본 보고서는 로컬 환경에서 실행되는 자동 분석 도구가 생성했습니다. "
        "외부 서버로 데이터가 전송되지 않습니다.",
        color=SUBTLE,
        size=9,
    )

    return rb.save(Path(output_path))


def _task_label(modeling: Optional[ModelingResult], config: AnalysisConfig) -> str:
    if modeling is None:
        return "탐색적 데이터 분석(EDA)"
    kind = "회귀 예측" if modeling.task_type == "regression" else "분류 예측"
    return f"{kind} — 목표 변수 '{modeling.target}'"


def default_report_name(prefix: str = "Analysis_Report") -> str:
    return f"{prefix}_{datetime.now():%Y%m%d_%H%M%S}.docx"
